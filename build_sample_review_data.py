"""
Build the synthetic study package the review module is tested against:

  data/sample_datasets/adsl.xpt          60 subjects, 3 arms (SAS transport v5)
  data/sample_datasets/adae.xpt          adverse events, TRTEMFL derived
  data/sample_outputs/t_14_1_1.rtf       Table 14.1.1 demographics (real RTF table)
  data/sample_outputs/t_14_3_1.docx      Table 14.3.1 TEAE overview (Word table)
  data/sample_planted.json               ground truth for the 3 planted problems

The first four subjects come from raw_data/demog.csv so the package ties back
to the existing sample data; the rest are seeded-random (same output every
run). Compared with data/sample_sap.pdf (built by build_sample_sap.py) the
package has exactly three deliberate inconsistencies:

  1. TEAE window   SAP 6.2: first dose to last dose + 7 days.
                   ADAE.TRTEMFL was derived with last dose + 30 days.
  2. Big N         SAP 4.5: Safety Population counts in every header.
                   Table 14.1.1 prints the ITT count for Placebo.
  3. Inventory     SAP 9 lists Table 14.3.2; it is not delivered.

Everything else is built to agree with the SAP (percentages, totals, etc.),
so a check that fires on anything but these three is a false positive.
All counts in the manifest are computed from the built data, not typed in.
"""

import datetime as dt
import json
import os
import random
from decimal import ROUND_HALF_UP, Decimal
from statistics import mean, median, stdev

import docx as docx_lib
import pandas as pd
import pyreadstat

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE, "data")
DS_DIR = os.path.join(DATA_DIR, "sample_datasets")
OUT_DIR = os.path.join(DATA_DIR, "sample_outputs")
MANIFEST = os.path.join(DATA_DIR, "sample_planted.json")

SEED = 20260919
STUDYID = "STUDY01"
ARMS = [("Placebo", 1), ("Drug A 50mg", 2), ("Drug A 100mg", 3)]
PER_ARM = 20

SAP_WINDOW_DAYS = 7       # SAP section 6.2
DATA_WINDOW_DAYS = 30     # what ADAE.TRTEMFL was (wrongly) derived with
N_LATE_AE = 9             # events inside the wrong window but outside the SAP one
N_VERY_LATE_AE = 3        # outside both windows (flag 'N' either way)

SAS_EPOCH = dt.date(1960, 1, 1)

RACE_LABELS = {
    "WHITE": "White",
    "BLACK OR AFRICAN AMERICAN": "Black or African American",
    "ASIAN": "Asian",
}
AE_TERMS = [
    ("Nervous system disorders", "Headache"),
    ("Nervous system disorders", "Dizziness"),
    ("Gastrointestinal disorders", "Nausea"),
    ("Gastrointestinal disorders", "Diarrhoea"),
    ("General disorders and administration site conditions", "Fatigue"),
    ("Skin and subcutaneous tissue disorders", "Rash"),
    ("Infections and infestations", "Nasopharyngitis"),
]
RELATED = ("POSSIBLE", "PROBABLE")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def rnd(x, places):
    """SAS-style rounding (half away from zero), returned as fixed text."""
    q = Decimal(1).scaleb(-places)
    return str(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))


def n_pct(n, big_n):
    return "0" if n == 0 else f"{n} ({rnd(100 * n / big_n, 1)}%)"


def sas_days(series):
    """datetime column -> SAS numeric date (days since 1960-01-01)."""
    return (series - pd.Timestamp(SAS_EPOCH)).dt.days.astype("float64")


# ---------------------------------------------------------------------------
# ADSL
# ---------------------------------------------------------------------------

def build_adsl(rng):
    raw = pd.read_csv(os.path.join(BASE, "raw_data", "demog.csv"))
    subjects = []
    for i, r in raw.iterrows():
        subjects.append({
            "SUBJID": f"{i + 1:03d}", "SITEID": int(r["SITE_NUM"]),
            "DOB": dt.date.fromisoformat(r["DOB"]), "SEX": r["GENDER"],
            "RACE": r["RACE_DESC"], "TRT01P": r["ARM_DESC"],
            "RANDDT": dt.date.fromisoformat(r["RANDOM_DATE"]),
        })

    remaining = []
    for arm, _ in ARMS:
        taken = sum(1 for s in subjects if s["TRT01P"] == arm)
        remaining += [arm] * (PER_ARM - taken)
    rng.shuffle(remaining)

    last_rand = subjects[-1]["RANDDT"]
    for arm in remaining:
        last_rand += dt.timedelta(days=rng.randint(1, 3))
        subjects.append({
            "SUBJID": f"{len(subjects) + 1:03d}", "SITEID": rng.choice([101, 102, 103]),
            "DOB": dt.date(1950, 1, 1) + dt.timedelta(days=rng.randint(0, 15000)),
            "SEX": rng.choice(["M", "F"]),
            "RACE": rng.choices(list(RACE_LABELS), weights=[6, 2, 2])[0],
            "TRT01P": arm, "RANDDT": last_rand,
        })

    # One Placebo subject (not one of the raw four) is randomised, never dosed.
    undosed = next(s["SUBJID"] for s in subjects[4:] if s["TRT01P"] == "Placebo")
    arm_n = dict(ARMS)
    rows = []
    for s in subjects:
        dosed = s["SUBJID"] != undosed
        trtsdt = s["RANDDT"] + dt.timedelta(days=rng.randint(3, 7)) if dosed else None
        trtedt = trtsdt + dt.timedelta(days=rng.randint(50, 60)) if dosed else None
        rows.append({
            "STUDYID": STUDYID, "USUBJID": f"{STUDYID}-{s['SUBJID']}", "SUBJID": s["SUBJID"],
            "SITEID": s["SITEID"],
            "AGE": int((s["RANDDT"] - s["DOB"]).days // 365.25),
            "SEX": s["SEX"], "RACE": s["RACE"],
            "TRT01P": s["TRT01P"], "TRT01A": s["TRT01P"] if dosed else "",
            "TRT01AN": float(arm_n[s["TRT01P"]]) if dosed else float("nan"),
            "ITTFL": "Y", "SAFFL": "Y" if dosed else "N",
            "TRTSDT": pd.Timestamp(trtsdt) if dosed else pd.NaT,
            "TRTEDT": pd.Timestamp(trtedt) if dosed else pd.NaT,
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# ADAE
# ---------------------------------------------------------------------------

def build_adae(rng, adsl):
    dosed = adsl[adsl["SAFFL"] == "Y"].reset_index(drop=True)
    events = []   # (subject row, start offset in days from TRTSDT, kind)

    for _, s in dosed.iterrows():
        dur = (s["TRTEDT"] - s["TRTSDT"]).days
        for _ in range(rng.choices([0, 1, 2, 3], weights=[35, 35, 20, 10])[0]):
            events.append((s, rng.randint(0, dur + SAP_WINDOW_DAYS), "in_window"))

    in_window_subj = {ev[0]["USUBJID"] for ev in events}
    # Pre-dose events: not treatment-emergent under any window.
    for _, s in dosed.sample(n=5, random_state=SEED).iterrows():
        events.append((s, -rng.randint(1, 10), "pre_dose"))

    # Planted problem 1: N_LATE_AE events starting 8-30 days after last dose.
    # Six go to subjects with no other TEAE, so the subject counts move too.
    quiet = [s for _, s in dosed.iterrows() if s["USUBJID"] not in in_window_subj]
    busy = [s for _, s in dosed.iterrows() if s["USUBJID"] in in_window_subj]
    late_subjects = quiet[:6] + rng.sample(busy, N_LATE_AE - 6)
    for s in late_subjects:
        dur = (s["TRTEDT"] - s["TRTSDT"]).days
        events.append((s, dur + rng.randint(SAP_WINDOW_DAYS + 1, DATA_WINDOW_DAYS), "late"))
    for s in rng.sample(busy, N_VERY_LATE_AE):
        dur = (s["TRTEDT"] - s["TRTSDT"]).days
        events.append((s, dur + rng.randint(DATA_WINDOW_DAYS + 5, DATA_WINDOW_DAYS + 15), "very_late"))

    rows = []
    for s, offset, _ in events:
        soc, pt = rng.choice(AE_TERMS)
        astdt = s["TRTSDT"] + pd.Timedelta(days=offset)
        rows.append({
            "STUDYID": STUDYID, "USUBJID": s["USUBJID"], "TRTA": s["TRT01A"],
            "AETERM": pt.lower(), "AEDECOD": pt, "AEBODSYS": soc,
            "AESEV": rng.choices(["MILD", "MODERATE", "SEVERE"], weights=[60, 30, 10])[0],
            "AESER": rng.choices(["N", "Y"], weights=[95, 5])[0],
            "AEREL": rng.choice(["NOT RELATED", "UNLIKELY", "POSSIBLE", "PROBABLE"]),
            "ASTDT": astdt,
            "AENDT": astdt + pd.Timedelta(days=rng.randint(0, 10)),
        })
    ae = pd.DataFrame(rows).sort_values(["USUBJID", "ASTDT"]).reset_index(drop=True)
    ae["AESEQ"] = ae.groupby("USUBJID").cumcount().astype("float64") + 1

    # The data's own TRTEMFL rule - last dose + DATA_WINDOW_DAYS. This is the bug.
    adsl_i = adsl.set_index("USUBJID")
    trtsdt = ae["USUBJID"].map(adsl_i["TRTSDT"])
    trtedt = ae["USUBJID"].map(adsl_i["TRTEDT"])
    emergent = (ae["ASTDT"] >= trtsdt) & (ae["ASTDT"] <= trtedt + pd.Timedelta(days=DATA_WINDOW_DAYS))
    ae["TRTEMFL"] = emergent.map({True: "Y", False: "N"})
    cols = ["STUDYID", "USUBJID", "AESEQ", "TRTA", "AETERM", "AEDECOD", "AEBODSYS",
            "AESEV", "AESER", "AEREL", "ASTDT", "AENDT", "TRTEMFL"]
    return ae[cols]


# ---------------------------------------------------------------------------
# Write .xpt
# ---------------------------------------------------------------------------

ADSL_LABELS = {
    "STUDYID": "Study Identifier", "USUBJID": "Unique Subject Identifier",
    "SUBJID": "Subject Identifier for the Study", "SITEID": "Study Site Identifier",
    "AGE": "Age", "SEX": "Sex", "RACE": "Race",
    "TRT01P": "Planned Treatment for Period 01", "TRT01A": "Actual Treatment for Period 01",
    "TRT01AN": "Actual Treatment for Period 01 (N)",
    "ITTFL": "Intent-To-Treat Population Flag", "SAFFL": "Safety Population Flag",
    "TRTSDT": "Date of First Exposure to Treatment", "TRTEDT": "Date of Last Exposure to Treatment",
}
ADAE_LABELS = {
    "STUDYID": "Study Identifier", "USUBJID": "Unique Subject Identifier",
    "AESEQ": "Sequence Number", "TRTA": "Actual Treatment",
    "AETERM": "Reported Term for the Adverse Event",
    "AEDECOD": "Dictionary-Derived Term", "AEBODSYS": "Body System or Organ Class",
    "AESEV": "Severity/Intensity", "AESER": "Serious Event", "AEREL": "Causality",
    "ASTDT": "Analysis Start Date", "AENDT": "Analysis End Date",
    "TRTEMFL": "Treatment Emergent Analysis Flag",
}


def write_xpt(df, name, labels, path):
    out = df.copy()
    formats = {}
    for col in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[col]):
            out[col] = sas_days(out[col])
            formats[col] = "DATE9"
    pyreadstat.write_xport(
        out, path, file_label=labels.get("_file", name), table_name=name,
        column_labels=[labels[c] for c in out.columns], file_format_version=5,
        variable_format=formats,
    )


# ---------------------------------------------------------------------------
# Tables (built from the data, so they agree with it except where planted)
# ---------------------------------------------------------------------------

def demog_table(adsl):
    saf = adsl[adsl["SAFFL"] == "Y"]
    groups = [(arm, saf[saf["TRT01A"] == arm]) for arm, _ in ARMS] + [("Total", saf)]

    # Planted problem 2: the Placebo header takes the ITT count, not SAFFL.
    header_n = {arm: len(g) for arm, g in groups}
    header_n["Placebo"] = int(((adsl["TRT01P"] == "Placebo") & (adsl["ITTFL"] == "Y")).sum())
    header = ["Parameter"] + [f"{arm} (N={header_n[arm]})" for arm, _ in groups]

    def line(label, fn):
        return [label] + [fn(g) for _, g in groups]

    def cat(label, col, value):
        # Denominator is each column's real Safety N - only the header is wrong.
        return line(f"  {label}", lambda g: n_pct(int((g[col] == value).sum()), len(g)))

    rows = [
        line("Age (years)", lambda g: ""),
        line("  n", lambda g: str(len(g))),
        line("  Mean (SD)", lambda g: f"{rnd(mean(g['AGE']), 1)} ({rnd(stdev(g['AGE']), 2)})"),
        line("  Median", lambda g: rnd(median(g["AGE"]), 1)),
        line("  Min, Max", lambda g: f"{int(g['AGE'].min())}, {int(g['AGE'].max())}"),
        line("Sex, n (%)", lambda g: ""),
        cat("Male", "SEX", "M"), cat("Female", "SEX", "F"),
        line("Race, n (%)", lambda g: ""),
    ] + [cat(label, "RACE", code) for code, label in RACE_LABELS.items()]
    return {
        "id": "14.1.1",
        "title": "Table 14.1.1 Summary of Demographic and Baseline Characteristics",
        "subtitle": "Safety Population",
        "header": header, "rows": rows,
        "footnotes": ["Percentages are based on the number of subjects in the Safety Population.",
                      "Source: ADSL. Program: t_14_1_1.sas"],
    }


def ae_table(adsl, adae):
    saf = adsl[adsl["SAFFL"] == "Y"]
    teae = adae[adae["TRTEMFL"] == "Y"]
    order = {"MILD": 1, "MODERATE": 2, "SEVERE": 3}

    groups = []
    for arm, _ in ARMS:
        groups.append((arm, saf[saf["TRT01A"] == arm]["USUBJID"]))
    groups.append(("Total", saf["USUBJID"]))
    header = ["Category"] + [f"{arm} (N={len(ids)})" for arm, ids in groups]

    def subjects(mask_df, ids):
        return mask_df[mask_df["USUBJID"].isin(ids)]["USUBJID"].nunique()

    max_sev = teae.assign(rank=teae["AESEV"].map(order)).groupby("USUBJID")["rank"].max()

    def line(label, fn):
        return [label] + [n_pct(fn(ids), len(ids)) for _, ids in groups]

    rows = [
        line("Any TEAE", lambda ids: subjects(teae, ids)),
        line("  Maximum severity: Mild", lambda ids: int((max_sev[max_sev.index.isin(ids)] == 1).sum())),
        line("  Maximum severity: Moderate", lambda ids: int((max_sev[max_sev.index.isin(ids)] == 2).sum())),
        line("  Maximum severity: Severe", lambda ids: int((max_sev[max_sev.index.isin(ids)] == 3).sum())),
        line("Any serious TEAE", lambda ids: subjects(teae[teae["AESER"] == "Y"], ids)),
        line("Any TEAE related to study drug", lambda ids: subjects(teae[teae["AEREL"].isin(RELATED)], ids)),
    ]
    return {
        "id": "14.3.1",
        "title": "Table 14.3.1 Overview of Treatment-Emergent Adverse Events",
        "subtitle": "Safety Population",
        "header": header, "rows": rows,
        "footnotes": ["Percentages are based on the number of subjects in the Safety Population.",
                      "A subject with more than one event is counted once per category.",
                      "Source: ADSL, ADAE. Program: t_14_3_1.sas"],
    }


# ---------------------------------------------------------------------------
# Writers: RTF (real \trowd table) and DOCX
# ---------------------------------------------------------------------------

def _rtf_escape(text):
    return text.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}")


def write_rtf(table, path):
    n_cols = len(table["header"])
    widths = [3200] + [1900] * (n_cols - 1)
    edges, x = [], 0
    for w in widths:
        x += w
        edges.append(x)
    cellx = "".join(f"\\cellx{e}" for e in edges)

    def row(cells, bold=False):
        b, b0 = ("\\b ", "\\b0 ") if bold else ("", "")
        body = "".join(f"\\intbl {b}{_rtf_escape(c)}{b0}\\cell" for c in cells)
        return f"\\trowd\\trgaph108{cellx}{body}\\row"

    parts = [
        "{\\rtf1\\ansi\\deff0", "{\\fonttbl{\\f0 Courier New;}}", "\\f0\\fs18",
        f"{_rtf_escape(table['title'])}\\par", f"{_rtf_escape(table['subtitle'])}\\par", "\\par",
        row(table["header"], bold=True),
    ] + [row(r) for r in table["rows"]] + ["\\par"] + [
        f"{_rtf_escape(f)}\\par" for f in table["footnotes"]
    ] + ["}"]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))


def write_docx(table, path):
    d = docx_lib.Document()
    d.add_paragraph(table["title"])
    d.add_paragraph(table["subtitle"])
    grid = d.add_table(rows=1 + len(table["rows"]), cols=len(table["header"]))
    for r, cells in enumerate([table["header"]] + table["rows"]):
        for c, val in enumerate(cells):
            grid.rows[r].cells[c].text = val
    for note in table["footnotes"]:
        d.add_paragraph(note)
    d.save(path)


# ---------------------------------------------------------------------------
# Manifest - ground truth, computed from the built data
# ---------------------------------------------------------------------------

def build_manifest(adsl, adae, demog):
    i = adsl.set_index("USUBJID")
    ae = adae.copy()
    ae["TRTSDT"] = ae["USUBJID"].map(i["TRTSDT"])
    ae["TRTEDT"] = ae["USUBJID"].map(i["TRTEDT"])
    sap_ok = (ae["ASTDT"] >= ae["TRTSDT"]) & (ae["ASTDT"] <= ae["TRTEDT"] + pd.Timedelta(days=SAP_WINDOW_DAYS))
    flagged = ae["TRTEMFL"] == "Y"
    late = ae[flagged & ~sap_ok]

    def subj_by_arm(mask):
        sub = ae[mask].merge(adsl[["USUBJID", "TRT01A"]], on="USUBJID")
        return {arm: int(sub[sub["TRT01A"] == arm]["USUBJID"].nunique()) for arm, _ in ARMS}

    saf_n = {arm: int(((adsl["SAFFL"] == "Y") & (adsl["TRT01A"] == arm)).sum()) for arm, _ in ARMS}
    itt_n = {arm: int(((adsl["ITTFL"] == "Y") & (adsl["TRT01P"] == arm)).sum()) for arm, _ in ARMS}
    printed = {h.split(" (N=")[0]: int(h.split("N=")[1].rstrip(")")) for h in demog["header"][1:]}
    sap_tables = ["14.1.1", "14.3.1", "14.3.2"]
    delivered = ["14.1.1", "14.3.1"]

    return {
        "note": "Ground truth for the planted inconsistencies. Computed by "
                "build_sample_review_data.py from the built datasets; the review "
                "checks must find these and report these counts.",
        "planted": [
            {
                "id": "P1", "check": "check_teae_window", "sap_ref": "SAP 6.2",
                "where": "ADAE.TRTEMFL",
                "summary": f"TRTEMFL derived with last dose + {DATA_WINDOW_DAYS} days; "
                           f"SAP says last dose + {SAP_WINDOW_DAYS} days.",
                "expected": {
                    "records_flagged_Y": int(flagged.sum()),
                    "flagged_Y_outside_sap_window": int(len(late)),
                    "subjects_with_any_teae_current": subj_by_arm(flagged),
                    "subjects_with_any_teae_under_sap": subj_by_arm(flagged & sap_ok),
                },
            },
            {
                "id": "P2", "check": "check_bign", "sap_ref": "SAP 4.5",
                "where": "Table 14.1.1 column header, Placebo",
                "summary": "Placebo header N is the ITT count; the table is a Safety Population table.",
                "expected": {
                    "printed_header_n": printed,
                    "adsl_saffl_n": saf_n,
                    "adsl_ittfl_n": itt_n,
                    "arms_that_disagree": [a for a in saf_n if printed[a] != saf_n[a]],
                },
            },
            {
                "id": "P3", "check": "check_tlf_inventory", "sap_ref": "SAP 9",
                "where": "Table inventory",
                "summary": "Table 14.3.2 is required by the SAP and was not delivered.",
                "expected": {"sap_tables": sap_tables, "delivered": delivered,
                             "missing": [t for t in sap_tables if t not in delivered]},
            },
        ],
    }


# ---------------------------------------------------------------------------

def build():
    os.makedirs(DS_DIR, exist_ok=True)
    os.makedirs(OUT_DIR, exist_ok=True)
    rng = random.Random(SEED)

    adsl = build_adsl(rng)
    adae = build_adae(rng, adsl)
    write_xpt(adsl, "ADSL", {**ADSL_LABELS, "_file": "Subject-Level Analysis Dataset"},
              os.path.join(DS_DIR, "adsl.xpt"))
    write_xpt(adae, "ADAE", {**ADAE_LABELS, "_file": "Adverse Events Analysis Dataset"},
              os.path.join(DS_DIR, "adae.xpt"))

    demog = demog_table(adsl)
    write_rtf(demog, os.path.join(OUT_DIR, "t_14_1_1.rtf"))
    write_docx(ae_table(adsl, adae), os.path.join(OUT_DIR, "t_14_3_1.docx"))

    manifest = build_manifest(adsl, adae, demog)
    assert manifest["planted"][0]["expected"]["flagged_Y_outside_sap_window"] == N_LATE_AE
    with open(MANIFEST, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2, sort_keys=True)
        f.write("\n")
    return manifest


if __name__ == "__main__":
    m = build()
    print("Wrote datasets, outputs and manifest under data/")
    for p in m["planted"]:
        print(" ", p["id"], p["check"], p["expected"])

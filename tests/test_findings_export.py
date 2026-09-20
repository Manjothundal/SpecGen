"""Piece C tests: review.findings - the Finding model, dedupe, root-cause
grouping and the Excel export (exact column order, round trip).

Run from the repo root:  python -m pytest tests/test_findings_export.py -v
"""

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from review.findings import (COLUMNS, Finding, dedupe, export_findings_xlsx, finalise,
                             order_findings, read_findings_xlsx)


def mk(output_file="Table 14.1.1 (t_14_1_1.rtf)", severity="major", rule_id="R-001", **over):
    base = {"output_file": output_file, "what_is_wrong": "Something does not agree.",
            "evidence": "1 of 4 columns differ", "source_ref": "SAP 4.5",
            "severity": severity, "rule_id": rule_id}
    base.update(over)
    return Finding.model_validate(base)


# --- the model -----------------------------------------------------------------

def test_the_seven_columns_are_exactly_as_agreed():
    assert COLUMNS == ["#", "Output / File", "What is wrong", "Evidence", "Source ref", "Severity", "Decision"]


def test_root_cause_id_is_nullable_and_defaults_to_none():
    f = mk()
    assert f.root_cause_id is None and f.decision == "open" and f.finding_id is None
    assert mk(root_cause_id="F-001").root_cause_id == "F-001"


@pytest.mark.parametrize("bad", [
    {"evidence": ""}, {"evidence": "   "}, {"evidence": "the numbers look wrong"},
    {"source_ref": ""}, {"source_ref": "  "}, {"output_file": ""}, {"what_is_wrong": ""},
    {"severity": "urgent"}, {"decision": "maybe"}, {"unknown_field": 1},
])
def test_a_finding_without_evidence_or_source_cannot_exist(bad):
    with pytest.raises(ValidationError):
        mk(**bad)


def test_a_decision_needs_a_person_and_a_time():
    with pytest.raises(ValidationError):
        mk(decision="accepted")
    with pytest.raises(ValidationError):
        mk(decision="rejected", decided_by="mhundal")
    f = mk(decision="accepted", decided_by="mhundal", decided_at="2026-09-19T11:00:00")
    assert f.decision == "accepted"


# --- dedupe ----------------------------------------------------------------------

def test_dedupe_same_rule_same_object_is_one_finding_at_the_highest_severity():
    a, b = mk(severity="minor"), mk(severity="critical", evidence="9 of 66 records")
    out = dedupe([a, b])
    assert len(out) == 1 and out[0].severity == "critical" and out[0].evidence == "1 of 4 columns differ"
    assert a.severity == "minor"                                  # the caller's object is untouched


def test_dedupe_keeps_different_objects_and_different_rules_and_hand_raised_findings():
    findings = [mk(), mk(output_file="Table 14.3.1 (t.rtf)"), mk(rule_id="R-002"),
                mk(rule_id=None), mk(rule_id=None)]
    assert len(dedupe(findings)) == 5


# --- ordering and root causes ------------------------------------------------------

def test_order_is_severity_then_output_then_rule():
    minor, major_b, major_a, critical = (mk(severity="minor"), mk(output_file="Table 14.3.1", severity="major"),
                                         mk(output_file="Table 14.1.1", severity="major"), mk(severity="critical"))
    assert order_findings([minor, major_b, major_a, critical]) == [critical, major_a, major_b, minor]


def test_a_root_cause_sits_directly_above_the_findings_that_point_to_it():
    root = mk("ADSL (adsl.xpt)", severity="minor", rule_id="R-010", finding_id="a")
    dep1 = mk("Table 14.3.1 (t.rtf)", severity="major", rule_id="R-002", finding_id="b", root_cause_id="a")
    dep2 = mk("Table 14.1.1 (t.rtf)", severity="major", rule_id="R-001", finding_id="c", root_cause_id="a")
    other = mk("Table 14.2.1 (t.rtf)", severity="major", rule_id="R-003", finding_id="d")
    order = order_findings([dep1, other, dep2, root])
    assert [f.finding_id for f in order] == ["a", "c", "b", "d"]   # root, its two findings, then the rest


def test_a_group_sorts_by_its_most_severe_member_but_the_root_stays_on_top():
    root = mk("ADSL (adsl.xpt)", severity="minor", finding_id="a")
    dep = mk("Table 14.1.1", severity="critical", rule_id="R-002", finding_id="b", root_cause_id="a")
    other = mk("Table 14.2.1", severity="major", rule_id="R-003", finding_id="c")
    assert [f.finding_id for f in order_findings([other, dep, root])] == ["a", "b", "c"]


def test_finalise_numbers_in_order_and_writes_the_link_text():
    criticals = [mk(f"Table 9.{i}.1", severity="critical", rule_id=f"R-{100 + i}") for i in range(11)]
    root = mk("ADaM/adsl.xpt", rule_id="R-010", finding_id="root", what_is_wrong="SAFFL is wrong in ADSL.")
    d1 = mk("Table 14.1.1 (t_14_1_1.rtf)", rule_id="R-001", finding_id="d1", root_cause_id="root")
    d2 = mk("Table 14.3.1.1 (t_14_3_1_1.rtf)", rule_id="R-002", finding_id="d2", root_cause_id="root")
    out = finalise(criticals + [d2, root, d1], run_id="RUN-1")

    assert [f.finding_id for f in out] == [f"F-{i:03d}" for i in range(1, 15)]
    by = {f.finding_id: f for f in out}
    cause = by["F-012"]
    assert cause.output_file == "ADaM/adsl.xpt" and cause.root_cause_id is None
    assert cause.what_is_wrong == ("SAFFL is wrong in ADSL. This is the dataset-level cause of the findings "
                                   "raised against Table 14.1.1 and Table 14.3.1.1.")
    pointing = [f for f in out if f.root_cause_id]
    assert [f.finding_id for f in pointing] == ["F-013", "F-014"]        # directly below the cause
    for f in pointing:
        assert f.root_cause_id == "F-012"
        assert f.what_is_wrong.endswith("Root cause: see finding F-012 against ADaM/adsl.xpt.")
    assert all(f.run_id == "RUN-1" for f in out)


def test_three_findings_are_joined_with_commas_and_and():
    root = mk("ADSL (adsl.xpt)", finding_id="r")
    deps = [mk(f"Table 14.{i}.1", rule_id=f"R-10{i}", finding_id=f"d{i}", root_cause_id="r") for i in (1, 2, 3)]
    out = finalise([root] + deps)
    assert "raised against Table 14.1.1, Table 14.2.1 and Table 14.3.1." in out[0].what_is_wrong


def test_finalise_is_repeatable_and_does_not_change_its_input():
    root = mk("ADSL (adsl.xpt)", finding_id="r")
    dep = mk("Table 14.1.1 (t.rtf)", rule_id="R-002", finding_id="d", root_cause_id="r")
    once = finalise([root, dep])
    twice = finalise(once)
    assert [f.model_dump() for f in once] == [f.model_dump() for f in twice]
    assert root.what_is_wrong == "Something does not agree." and root.finding_id == "r"


@pytest.mark.parametrize("make", [
    lambda: [mk(finding_id="a", root_cause_id="zzz")],                                    # root not in the run
    lambda: [mk(finding_id="a", root_cause_id="a")],                                      # points at itself
    lambda: [mk(finding_id="r"), mk(rule_id="R-2", finding_id="m", root_cause_id="r"),
             mk(rule_id="R-3", finding_id="d", root_cause_id="m")],                       # a chain
    lambda: [mk(finding_id="x"), mk(rule_id="R-2", finding_id="x")],                      # duplicate ids
])
def test_finalise_rejects_broken_links(make):
    with pytest.raises(ValueError):
        finalise(make())


def test_findings_with_no_links_are_left_alone():
    out = finalise([mk(), mk(rule_id="R-002")])
    assert [f.what_is_wrong for f in out] == ["Something does not agree."] * 2
    assert all(f.root_cause_id is None for f in out)


# --- the Excel export ----------------------------------------------------------------

@pytest.fixture
def linked():
    root = mk("ADaM/adsl.xpt", severity="critical", rule_id="R-010", finding_id="r")
    d1 = mk("Table 14.1.1 (t.rtf)", rule_id="R-001", finding_id="d1", root_cause_id="r")
    d2 = mk("Table 14.3.1 (t.rtf)", severity="minor", rule_id="R-002", finding_id="d2", root_cause_id="r",
            decision="accepted", decided_by="mhundal", decided_at="2026-09-19T11:00:00")
    other = mk("Table 14.2.1 (t.rtf)", rule_id="R-003", finding_id="o", evidence="3 of 9 rows")
    return finalise([other, d2, root, d1])


def test_export_has_exactly_the_seven_columns_and_no_root_cause_column(tmp_path, linked):
    path = export_findings_xlsx(linked, str(tmp_path / "f.xlsx"))
    ws = load_workbook(path)["Findings"]
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.max_column == 7
    assert not any("root" in str(c.value).lower() for c in ws[1])


def test_export_puts_a_root_cause_above_the_findings_that_point_to_it(tmp_path, linked):
    rows = read_findings_xlsx(export_findings_xlsx(linked, str(tmp_path / "f.xlsx")))
    assert [r["Output / File"] for r in rows] == ["ADaM/adsl.xpt", "Table 14.1.1 (t.rtf)",
                                                  "Table 14.3.1 (t.rtf)", "Table 14.2.1 (t.rtf)"]
    assert "This is the dataset-level cause of the findings raised against Table 14.1.1 and Table 14.3.1." \
        in rows[0]["What is wrong"]
    assert all("Root cause: see finding F-001 against ADaM/adsl.xpt." in r["What is wrong"] for r in rows[1:3])
    assert "Root cause" not in rows[3]["What is wrong"]


def test_export_sorts_by_itself_even_when_given_findings_out_of_order(tmp_path, linked):
    rows = read_findings_xlsx(export_findings_xlsx(list(reversed(linked)), str(tmp_path / "f.xlsx")))
    assert [r["#"] for r in rows] == [1, 2, 3, 4]


def test_the_hash_column_is_the_number_in_the_finding_id_even_for_a_filtered_export(tmp_path, linked):
    subset = [f for f in linked if f.finding_id in ("F-002", "F-004")]
    rows = read_findings_xlsx(export_findings_xlsx(subset, str(tmp_path / "f.xlsx")))
    assert [r["#"] for r in rows] == [2, 4]


def test_the_hash_column_counts_from_one_when_findings_have_no_ids(tmp_path):
    rows = read_findings_xlsx(export_findings_xlsx([mk(), mk(rule_id="R-2")], str(tmp_path / "f.xlsx")))
    assert [r["#"] for r in rows] == [1, 2]


def test_round_trip_gives_back_every_value(tmp_path, linked):
    rows = read_findings_xlsx(export_findings_xlsx(linked, str(tmp_path / "f.xlsx")))
    assert len(rows) == len(linked) == 4
    for row, f in zip(rows, order_findings(linked)):
        assert row == {"#": int(f.finding_id[2:]), "Output / File": f.output_file,
                       "What is wrong": f.what_is_wrong, "Evidence": f.evidence, "Source ref": f.source_ref,
                       "Severity": f.severity.capitalize(), "Decision": f.decision.capitalize()}
    assert {r["Decision"] for r in rows} == {"Open", "Accepted"}
    assert [r["Severity"] for r in rows][0] == "Critical"


def test_empty_export_is_a_header_only_sheet(tmp_path):
    ws = load_workbook(export_findings_xlsx([], str(tmp_path / "f.xlsx")))["Findings"]
    assert ws.max_row == 1 and ws.auto_filter.ref == "A1:G1"
    assert read_findings_xlsx(str(tmp_path / "f.xlsx")) == []


def test_export_uses_the_house_header_style(tmp_path, linked):
    ws = load_workbook(export_findings_xlsx(linked, str(tmp_path / "f.xlsx")))["Findings"]
    head = ws["A1"]
    assert head.font.bold and head.font.name == "Arial" and head.font.color.rgb.endswith("FFFFFF")
    assert head.fill.start_color.rgb.endswith("1A3C6E")
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:G5"
    assert ws["F2"].fill.start_color.rgb.endswith("FFC7CE")             # critical is red


def test_run_info_sheet_is_second_and_optional(tmp_path, linked):
    plain = load_workbook(export_findings_xlsx(linked, str(tmp_path / "a.xlsx")))
    assert plain.sheetnames == ["Findings"]
    info = {"summary": [("Draft rules", "INCLUDED")],
            "rules": [{"rule_id": "R-001", "version": 1, "status": "draft", "check": "check_bign",
                       "sap_ref": "SAP 4.5", "result": "1 finding(s)"}]}
    wb = load_workbook(export_findings_xlsx(linked, str(tmp_path / "b.xlsx"), run_info=info))
    assert wb.sheetnames == ["Findings", "Run info"]
    cells = [c for row in wb["Run info"].iter_rows(values_only=True) for c in row if c is not None]
    assert "INCLUDED" in cells and "check_bign" in cells and "draft" in cells


def test_reading_a_sheet_with_the_wrong_columns_is_refused(tmp_path):
    wb = Workbook()
    wb.active.title = "Findings"
    wb.active.append(["#", "Output / File", "What is wrong", "Evidence", "Source ref", "Severity", "Decision", "Root cause"])
    wb.save(str(tmp_path / "bad.xlsx"))
    with pytest.raises(ValueError):
        read_findings_xlsx(str(tmp_path / "bad.xlsx"))
